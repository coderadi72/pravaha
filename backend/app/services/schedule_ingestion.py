"""CSV/XLSX normalization, preview validation and immutable schedule-version diffs."""
import csv,io,math,re,json
from hashlib import sha256
from datetime import date,datetime
from openpyxl import load_workbook
from backend.app.core.errors import ApiError
from backend.app.services.document_processor import decode_file

FIELDS={"activityId","activityName","wbs","projectId","discipline","location","plannedStart","plannedEnd","actualStart","actualEnd","plannedDuration","actualDuration","progress","predecessors","successors","level","parentActivityId","teamId","totalFloat"}
ALIASES={"activityid":"activityId","activityname":"activityName","name":"activityName","project":"projectId","projectid":"projectId","plannedfinish":"plannedEnd","plannedend":"plannedEnd","plannedstart":"plannedStart","actualstart":"actualStart","actualfinish":"actualEnd","actualend":"actualEnd","percentcomplete":"progress","percentcomplete":"progress","progress":"progress","activitylevel":"level","responsibleteam":"teamId","teamid":"teamId","planned duration":"plannedDuration"}
for _field in FIELDS:ALIASES[re.sub(r'[^a-z0-9]','',_field.lower())]=_field
DEFAULT_MAX_ROWS=2000
DEFAULT_MAX_COLUMNS=60
DEFAULT_MAX_RELATIONSHIPS=10000
DEFAULT_DEPENDENCY_CELL_MAX_COUNT=50
DEFAULT_DEPENDENCY_MAX_LAG_DAYS=3650


def relationship_list(value, max_count=DEFAULT_DEPENDENCY_CELL_MAX_COUNT, max_lag=DEFAULT_DEPENDENCY_MAX_LAG_DAYS):
    if not value:return []
    result=[]
    if len(str(value).split(','))>max_count:raise ValueError(f'At most {max_count} dependencies per activity column.')
    for part in str(value).split(','):
        pieces=[p.strip() for p in part.strip().split(':')]
        if not pieces[0] or len(pieces)>3:raise ValueError('Invalid dependency; use ActivityID:FS:lagDays (comma separated).')
        relation=pieces[1].upper() if len(pieces)>1 else 'FS'
        lag=float(pieces[2]) if len(pieces)>2 else 0.0
        if relation not in {'FS','SS','FF','SF'} or not math.isfinite(lag) or abs(lag)>max_lag:raise ValueError('Invalid dependency type or lag.')
        result.append({'activityId':pieces[0],'type':relation,'lagDays':lag})
    return result


def parse_date(value):
    if isinstance(value,datetime):return value.date().isoformat()
    if isinstance(value,date):return value.isoformat()
    try:return date.fromisoformat(str(value)).isoformat()
    except ValueError:raise ValueError('Dates must use YYYY-MM-DD or native Excel dates.') from None


def graph_valid(ids,edges):
    successors={i:[] for i in ids};indegree={i:0 for i in ids}
    for e in edges:
        if e['predecessor']==e['successor']:raise ValueError('Self dependencies are invalid.')
        if e['predecessor'] not in ids or e['successor'] not in ids:raise ValueError('Dependency refers to an activity absent from this imported schedule.')
        successors[e['predecessor']].append(e['successor']);indegree[e['successor']]+=1
    stack=[i for i in ids if not indegree[i]];visited=0
    while stack:
        node=stack.pop();visited+=1
        for target in successors[node]:
            indegree[target]-=1
            if not indegree[target]:stack.append(target)
    if visited!=len(ids):raise ValueError('Dependency cycle detected.')


def preview(body,project_id,settings=None):
    max_rows=getattr(settings,'schedule_max_rows',DEFAULT_MAX_ROWS)
    max_columns=getattr(settings,'schedule_max_columns',DEFAULT_MAX_COLUMNS)
    max_relationships=getattr(settings,'schedule_max_relationships',DEFAULT_MAX_RELATIONSHIPS)
    dependency_cell_max_count=getattr(settings,'dependency_cell_max_count',DEFAULT_DEPENDENCY_CELL_MAX_COUNT)
    dependency_max_lag_days=getattr(settings,'dependency_max_lag_days',DEFAULT_DEPENDENCY_MAX_LAG_DAYS)
    file=decode_file(body,settings)
    if file['filename'].lower().endswith('.csv'):
        try:
            reader=csv.reader(io.StringIO(file['content'].decode('utf-8-sig')),strict=True)
            rows=[]
            for row in reader:
                rows.append(row)
                if len(rows)>max_rows+1:raise ApiError(413,'IMPORT_TOO_LARGE',f'Schedule imports are limited to {max_rows:,} rows.')
        except csv.Error:raise ApiError(400,'INVALID_SCHEDULE','Malformed CSV input.') from None
        source='csv'
    elif file['filename'].lower().endswith('.xlsx'):
        try:
            workbook=load_workbook(io.BytesIO(file['content']),read_only=True,data_only=False,keep_links=False)
            if len(workbook.worksheets)!=1:raise ApiError(400,'INVALID_SCHEDULE','Choose one worksheet; multi-sheet imports are not silently discarded.')
            sheet=workbook.active
            if sheet.max_row>max_rows+1 or sheet.max_column>max_columns:raise ApiError(413,'IMPORT_TOO_LARGE',f'Schedule imports allow {max_rows:,} rows and {max_columns} columns.')
            sheet.reset_dimensions()
            rows=[]
            try:
                for row in sheet.iter_rows(values_only=True):
                    if len(row)>max_columns or len(rows)>=max_rows+1:raise ApiError(413,'IMPORT_TOO_LARGE',f'Schedule imports allow {max_rows:,} rows and {max_columns} columns.')
                    rows.append(row)
            finally:workbook.close()
        except ApiError:raise
        except Exception:raise ApiError(400,'INVALID_SCHEDULE','Excel workbook could not be parsed safely.') from None
        source='xlsx'
    else:raise ApiError(400,'INVALID_SCHEDULE','Schedule files must be CSV or XLSX.')
    if not rows or len(rows)<2 or len(rows[0])>max_columns:raise ApiError(400,'INVALID_SCHEDULE',f'Provide a header and at least one activity, with at most {max_columns} columns.')
    headers=[str(h or '').strip() for h in rows[0]]
    if any(not h for h in headers) or len(set(headers))!=len(headers):raise ApiError(400,'INVALID_SCHEDULE','Headers must be nonempty and unique.')
    if any(k not in headers or v not in FIELDS and v != '' for k,v in body.mapping.items()):raise ApiError(400,'INVALID_MAPPING','Column mappings must reference existing headers and supported target fields.')
    mapping={h:body.mapping[h] if h in body.mapping else ALIASES.get(re.sub(r'[^a-z0-9]','',h.lower()), '') for h in headers}
    targets=[v for v in mapping.values() if v]
    if len(targets)!=len(set(targets)):raise ApiError(400,'INVALID_MAPPING','Multiple columns map to the same field.')
    required={'activityId','activityName','plannedStart','plannedEnd'}
    missing=required-set(targets)
    if missing:raise ApiError(400,'INVALID_MAPPING','Missing required mappings: '+', '.join(sorted(missing)))
    accepted,rejected,warnings,seen=[],[],[],set()
    unsupported=[h for h,v in mapping.items() if not v]
    if unsupported:warnings.append('Unsupported columns are preserved in extraColumns: '+', '.join(unsupported))
    for number,values in enumerate(rows[1:],start=2):
        if not any(v is not None and str(v).strip() for v in values):continue
        try:
            if len(values)>len(headers):raise ValueError('Row has more cells than headers.')
            original=dict(zip(headers,values));data={mapping[h]:v for h,v in original.items() if mapping[h] and v is not None and str(v).strip()}
            if any(str(v).lstrip().startswith('=') for v in values if isinstance(v,str)):raise ValueError('Formula cells are not accepted.')
            for field in required:
                if not data.get(field):raise ValueError('Missing '+field)
            identifier=str(data['activityId']).strip()
            if len(identifier)>120 or ':' in identifier or ',' in identifier:raise ValueError('Invalid or oversized activity ID.')
            if identifier in seen:raise ValueError('Duplicate activity ID: '+identifier)
            seen.add(identifier)
            for field in ['activityId','activityName','wbs','discipline','location','level','parentActivityId','teamId']:
                if field in data:data[field]=str(data[field]).strip()
                if len(data.get(field,''))>500:raise ValueError('Field exceeds 500 characters: '+field)
            if data.get('projectId') and str(data['projectId'])!=project_id:raise ValueError('Row belongs to another project.')
            data.setdefault('discipline','')
            data.setdefault('location','')
            data['projectId']=project_id
            for field in ['plannedStart','plannedEnd','actualStart','actualEnd']:
                if field in data:data[field]=parse_date(data[field])
            if data['plannedEnd']<data['plannedStart']:raise ValueError('Planned finish precedes planned start.')
            if data.get('actualStart') and data.get('actualEnd') and data['actualEnd']<data['actualStart']:raise ValueError('Actual finish precedes actual start.')
            for field in ['plannedDuration','actualDuration','progress','totalFloat']:
                if field in data:
                    value=float(data[field])
                    if not math.isfinite(value) or (field!='totalFloat' and value<0) or (field=='progress' and value>100):raise ValueError('Invalid numeric value: '+field)
                    data[field]=value
            data['level']=data.get('level','L6').upper()
            if data['level'] not in {f'L{i}' for i in range(7)}:raise ValueError('Level must be L0 through L6.')
            data['predecessors']=relationship_list(data.get('predecessors'),dependency_cell_max_count,dependency_max_lag_days)
            data['successors']=relationship_list(data.get('successors'),dependency_cell_max_count,dependency_max_lag_days)
            data['extraColumns']={h:str(original.get(h,'')) for h in unsupported}
            data['sourceRow']=number
            accepted.append(data)
        except (ValueError,TypeError,OverflowError) as error:rejected.append({'row':number,'errors':[str(error)]})
    edges={}
    for row in accepted:
        for key in ['predecessors','successors']:
            for dependency in row[key]:
                pred,succ=(dependency['activityId'],row['activityId']) if key=='predecessors' else (row['activityId'],dependency['activityId'])
                edge={'predecessor':pred,'successor':succ,'type':dependency['type'],'lagDays':dependency['lagDays']}
                identity=(pred,succ,edge['type'])
                if identity in edges and edges[identity]!=edge:rejected.append({'row':row['sourceRow'],'errors':['Conflicting dependency lag declarations.']})
                edges[identity]=edge
    if len(edges)>max_relationships:raise ApiError(413,'IMPORT_TOO_LARGE',f'At most {max_relationships:,} schedule relationships are supported.')
    try:
        graph_valid({r['activityId'] for r in accepted},list(edges.values()))
        for row in accepted:
            if row.get('parentActivityId') and row['parentActivityId'] not in {r['activityId'] for r in accepted}:raise ValueError('Hierarchy parent is absent from schedule.')
        graph_valid({r['activityId'] for r in accepted}, [{'predecessor':r['parentActivityId'],'successor':r['activityId']} for r in accepted if r.get('parentActivityId')])
    except ValueError as error:rejected.append({'row':None,'errors':[str(error)]})
    preview_checksum=sha256((file['checksum']+json.dumps(mapping,sort_keys=True)).encode()).hexdigest()
    return {'filename':file['filename'],'sourceFormat':source,'fileChecksum':file['checksum'],'checksum':preview_checksum,'headers':headers,'mapping':mapping,'rows':accepted,'dependencies':list(edges.values()),'accepted':len(accepted),'rejected':len(rejected),'errors':rejected,'warnings':warnings,'canImport':bool(accepted) and not rejected}


def compare_versions(before,after):
    old={r['activityId']:r for r in before};new={r['activityId']:r for r in after}
    keys=['activityName','plannedStart','plannedEnd','actualStart','actualEnd','plannedDuration','actualDuration','progress','predecessors','successors','teamId','parentActivityId','totalFloat']
    return {'added':sorted(new.keys()-old.keys()),'removed':sorted(old.keys()-new.keys()),'changed':[{'activityId':identifier,'fields':{k:{'previous':old[identifier].get(k),'new':new[identifier].get(k)} for k in keys if old[identifier].get(k)!=new[identifier].get(k)}} for identifier in sorted(old.keys()&new.keys()) if any(old[identifier].get(k)!=new[identifier].get(k) for k in keys)]}
